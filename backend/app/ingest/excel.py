"""Read a requirements spreadsheet with loosely named columns (EN/VI)."""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HEADER_ALIASES: dict[str, list[str]] = {
    'ref': ['id', 'req id', 'requirement id', 'req', 'key', 'mã', 'ma', 'mã yêu cầu', 'no', 'stt', '#'],
    'feature': ['feature', 'module', 'epic', 'title', 'chức năng', 'tính năng', 'tiêu đề', 'màn hình'],
    'text': [
        'requirement', 'requirements', 'description', 'user story', 'story', 'detail', 'details',
        'mô tả', 'yêu cầu', 'nội dung',
    ],
    'acceptance': [
        'acceptance criteria', 'acceptance', 'ac', 'expected', 'expected result',
        'tiêu chí', 'tiêu chí chấp nhận', 'kết quả mong đợi',
    ],
    'priority': ['priority', 'độ ưu tiên', 'ưu tiên', 'severity'],
}


def _norm(v: Any) -> str:
    return str(v).strip() if v is not None else ''


def _match_columns(header: list[str]) -> dict[str, int]:
    cols: dict[str, int] = {}
    lowered = [h.lower() for h in header]
    for field, aliases in HEADER_ALIASES.items():
        for i, h in enumerate(lowered):
            if h in aliases and i not in cols.values():
                cols[field] = i
                break
    return cols


def parse_requirements(path: str | Path) -> list[dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[dict[str, str]] = []
    for ws in wb.worksheets:
        rows = [[_norm(c) for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        # header = first row with at least 2 non-empty cells
        h_idx = next((i for i, r in enumerate(rows) if sum(1 for c in r if c) >= 2), 0)
        header = rows[h_idx]
        cols = _match_columns(header)
        for n, row in enumerate(rows[h_idx + 1 :], start=1):
            get = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else ''  # noqa: E731
            text = get('text')
            if not text:
                # Unknown layout: keep the whole row as "Header: value" pairs.
                text = '; '.join(f'{header[i] or i}: {c}' for i, c in enumerate(row) if c)
            if not text:
                continue
            out.append({
                'ref': get('ref') or f'{ws.title}-{n}',
                'feature': get('feature'),
                'text': text,
                'acceptance': get('acceptance'),
                'priority': get('priority'),
            })
    wb.close()
    return out
